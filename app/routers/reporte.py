# app/routers/reporte.py
from datetime import date
from decimal import Decimal
from sqlalchemy import text

from io import BytesIO

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


from collections import defaultdict
from openpyxl.styles import Border, Side
from openpyxl.worksheet.properties import PageSetupProperties
from app.deps import require_permission

from fastapi import APIRouter, Depends
from sqlalchemy import and_, func, select, true
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import (
    Bodega,
    Categoria,
    Compra,
    CompraDetalle,
    Movimiento,
    MovimientoDetalle,
    Producto,
    ProductoPresentacion,
    Saldo,
    UnidadMedida,
)

router = APIRouter(prefix="/reportes", tags=["Reportes"])


def obtener_filas_stock(db: Session) -> list[dict]:
    # Bloque 1: ENTRADAS. Por producto y bodega: primera/última compra y totales
        # Bloque 1: ENTRADAS (se leen de las compras, igual que tu pantalla de inventario)
    entradas = (
        select(
            ProductoPresentacion.producto_id.label("producto_id"),
            Movimiento.bodega_destino_id.label("bodega_id"),
            func.min(Compra.fecha_compra).label("primera_entrada"),
            func.max(Compra.fecha_compra).label("ultima_entrada"),
            func.sum(CompraDetalle.cantidad_base).label("total_ingresado"),
            func.sum(CompraDetalle.cantidad_base * CompraDetalle.costo_unitario_base).label("costo_total"),
        )
        .select_from(CompraDetalle)
        .join(Compra, Compra.id == CompraDetalle.compra_id)
        .join(Movimiento, Movimiento.compra_id == Compra.id)
        .join(ProductoPresentacion, ProductoPresentacion.id == CompraDetalle.presentacion_id)
        .where(Movimiento.bodega_destino_id.is_not(None))
        .group_by(ProductoPresentacion.producto_id, Movimiento.bodega_destino_id)
        .subquery()
    )

    # Bloque 2: SALIDAS. Por producto y bodega: última salida y total salido
    salidas = (
        select(
            MovimientoDetalle.producto_id.label("producto_id"),
            Movimiento.bodega_origen_id.label("bodega_id"),
            func.max(Movimiento.fecha).label("ultima_salida"),
            func.sum(MovimientoDetalle.cantidad).label("total_salido"),
        )
        .join(Movimiento, Movimiento.id == MovimientoDetalle.movimiento_id)
        .where(Movimiento.bodega_origen_id.is_not(None))
        .group_by(MovimientoDetalle.producto_id, Movimiento.bodega_origen_id)
        .subquery()
    )

    # Bloque 3: TODOS los productos en TODAS las bodegas, y pegamos el saldo,
    # las entradas y las salidas si existen (outerjoin = si no hay, queda vacío)
    consulta = (
        select(
            Producto.codigo,
            Producto.nombre.label("producto"),
            Categoria.nombre.label("categoria"),
            UnidadMedida.abreviatura.label("unidad"),
            Bodega.nombre.label("bodega"),
            func.coalesce(Saldo.cantidad, 0).label("stock"),
            entradas.c.primera_entrada,
            entradas.c.ultima_entrada,
            entradas.c.total_ingresado,
            entradas.c.costo_total,
            salidas.c.ultima_salida,
            salidas.c.total_salido,
        )
        .select_from(Producto)
        .join(Categoria, Categoria.id == Producto.categoria_id)
        .join(UnidadMedida, UnidadMedida.id == Producto.unidad_base_id)
        .join(Bodega, true())
        .outerjoin(Saldo, and_(Saldo.producto_id == Producto.id, Saldo.bodega_id == Bodega.id))
        .outerjoin(entradas, and_(entradas.c.producto_id == Producto.id, entradas.c.bodega_id == Bodega.id))
        .outerjoin(salidas, and_(salidas.c.producto_id == Producto.id, salidas.c.bodega_id == Bodega.id))
        .order_by(Bodega.nombre, Producto.nombre)
    )

    hoy = date.today()
    filas = []
    for r in db.execute(consulta).all():
        stock = Decimal(r.stock)
        ingresado = r.total_ingresado or Decimal("0")
        costo_promedio = (r.costo_total / ingresado) if ingresado else Decimal("0")
        ultima_salida = r.ultima_salida.date() if r.ultima_salida else None

        filas.append({
            "codigo": r.codigo,
            "producto": r.producto,
            "categoria": r.categoria,
            "bodega": r.bodega,
            "unidad": r.unidad,
            "stock": stock,
            "total_ingresado": ingresado,
            "total_salido": r.total_salido or Decimal("0"),
            "costo_promedio": round(costo_promedio, 4),
            "valor_total": round(stock * costo_promedio, 2),
            "primera_entrada": r.primera_entrada,
            "ultima_entrada": r.ultima_entrada,
            "ultima_salida": ultima_salida,
            "dias_en_bodega": (hoy - r.ultima_entrada).days if r.ultima_entrada else None,
            "dias_sin_salir": (hoy - ultima_salida).days if ultima_salida else None,
            "estado": "Con stock" if stock > 0 else "Sin stock",
        })
    return filas


# Endpoint temporal para revisar los datos antes de armar el Excel
@router.get("/stock")
def stock_json(db: Session = Depends(get_db)):
    return obtener_filas_stock(db)

#Reporte de Stock en Excel

@router.get("/stock/excel")
def stock_excel(db: Session = Depends(get_db)):
    filas = obtener_filas_stock(db)

    # Control: debe dar 0. Si no, hay movimientos mal registrados
    for f in filas:
        f["diferencia"] = f["stock"] - (f["total_ingresado"] - f["total_salido"])

    libro = Workbook()

    # ================== HOJA 1: RESUMEN ==================
    resumen = libro.active
    resumen.title = "Resumen"

    con_stock = [f for f in filas if f["estado"] == "Con stock"]
    sin_stock = [f for f in filas if f["estado"] == "Sin stock"]
    lentos = [f for f in con_stock if f["dias_en_bodega"] is not None and f["dias_en_bodega"] > 90]
    con_diferencia = [f for f in filas if f["diferencia"] != 0]

    datos_resumen = [
        ("Fecha del reporte", date.today(), "dd/mm/yyyy"),
        ("Líneas producto-bodega", len(filas), "0"),
        ("Con stock", len(con_stock), "0"),
        ("Sin stock", len(sin_stock), "0"),
        ("Valor total del inventario", sum(f["valor_total"] for f in filas), "$#,##0.00"),
        ("Con más de 90 días en bodega", len(lentos), "0"),
        ("Con diferencia en el control", len(con_diferencia), "0"),
    ]

    resumen["A1"] = "REPORTE DE STOCK"
    resumen["A1"].font = Font(bold=True, size=16)
    for i, (titulo, valor, formato) in enumerate(datos_resumen, start=3):
        resumen.cell(row=i, column=1, value=titulo).font = Font(bold=True)
        celda = resumen.cell(row=i, column=2, value=valor)
        celda.number_format = formato
        celda.alignment = Alignment(horizontal="left")
    resumen.column_dimensions["A"].width = 34
    resumen.column_dimensions["B"].width = 18

    # ================== HOJA 2: STOCK DETALLADO ==================
    hoja = libro.create_sheet("Stock detallado")

    columnas = [
        ("Código", "codigo", 16, None),
        ("Producto", "producto", 42, None),
        ("Categoría", "categoria", 20, None),
        ("Bodega", "bodega", 20, None),
        ("Unidad", "unidad", 9, None),
        ("Stock actual", "stock", 13, "#,##0.000"),
        ("Total ingresado", "total_ingresado", 15, "#,##0.000"),
        ("Total salido", "total_salido", 13, "#,##0.000"),
        ("Diferencia", "diferencia", 12, "#,##0.000"),
        ("Costo promedio", "costo_promedio", 15, "$#,##0.0000"),
        ("Valor total", "valor_total", 14, "$#,##0.00"),
        ("Primera entrada", "primera_entrada", 16, "dd/mm/yyyy"),
        ("Última entrada", "ultima_entrada", 16, "dd/mm/yyyy"),
        ("Última salida", "ultima_salida", 16, "dd/mm/yyyy"),
        ("Días en bodega", "dias_en_bodega", 15, "0"),
        ("Días sin salir", "dias_sin_salir", 15, "0"),
        ("Estado", "estado", 12, None),
    ]

    relleno_titulo = PatternFill("solid", fgColor="1F6FEB")
    relleno_sin_stock = PatternFill("solid", fgColor="FDE8E8")
    relleno_lento = PatternFill("solid", fgColor="FFF4CC")

    # Encabezados
    for col, (titulo, clave, ancho, formato) in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=col, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno_titulo
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(col)].width = ancho

    # Filas de datos
    for num_fila, f in enumerate(filas, start=2):
        for col, (titulo, clave, ancho, formato) in enumerate(columnas, start=1):
            valor = f[clave]
            # Producto con stock pero sin compra registrada
            if clave in ("primera_entrada", "ultima_entrada") and valor is None and f["estado"] == "Con stock":
                valor = "Sin compra registrada"
            celda = hoja.cell(row=num_fila, column=col, value=valor)
            if formato:
                celda.number_format = formato

            if f["estado"] == "Sin stock":
                celda.fill = relleno_sin_stock
            elif clave == "dias_en_bodega" and valor is not None and valor > 90:
                celda.fill = relleno_lento

    hoja.freeze_panes = "A2"          # la fila de títulos queda fija al bajar
    hoja.auto_filter.ref = hoja.dimensions  # filtros en los encabezados

    # ================== ENTREGAR EL ARCHIVO ==================
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"stock_{date.today():%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
    
    
    # =====================================================================
# NUEVO: Reporte de stock por bodega (una hoja por bodega)
# Ruta: GET /reportes/stock-bodegas/excel
# =====================================================================
def _nombre_hoja_excel(nombre: str, usados: set) -> str:
    # Excel no acepta [ ] : * ? / \ y limita el nombre a 31 caracteres
    limpio = "".join(c for c in nombre if c not in '[]:*?/\\')[:31] or "Bodega"
    base, n = limpio, 2
    while limpio in usados:
        sufijo = f" ({n})"
        limpio = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(limpio)
    return limpio


@router.get("/stock-bodegas/excel")
def stock_por_bodega_excel(
    db: Session = Depends(get_db),
    usuario=Depends(require_permission("inventario.ver")),
):
    solo_con_stock = True  # False = también lista los productos con stock 0

    bodegas = db.scalars(select(Bodega).order_by(Bodega.nombre)).all()

    consulta = (
        select(
            Saldo.bodega_id.label("bodega_id"),
            Producto.codigo,
            Producto.nombre.label("producto"),
            UnidadMedida.abreviatura.label("presentacion"),
            Saldo.cantidad.label("stock"),
        )
        .select_from(Saldo)
        .join(Producto, Producto.id == Saldo.producto_id)
        .join(UnidadMedida, UnidadMedida.id == Producto.unidad_base_id)
        .order_by(Producto.nombre)
    )
    if solo_con_stock:
        consulta = consulta.where(Saldo.cantidad > 0)

    filas_por_bodega = defaultdict(list)
    for r in db.execute(consulta).all():
        filas_por_bodega[r.bodega_id].append(r)

    azul = "1F6FEB"
    relleno_titulo = PatternFill("solid", fgColor=azul)
    relleno_alterno = PatternFill("solid", fgColor="F3F7FF")
    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    columnas = [
        ("CÓDIGO DE PRODUCTO", 22, None, "left"),
        ("NOMBRE", 52, None, "left"),
        ("TIPO DE PRESENTACIÓN", 24, None, "center"),
        ("CANTIDAD (STOCK DISPONIBLE)", 28, "#,##0.000", "right"),
    ]
    fila_titulos = 6

    libro = Workbook()
    nombres_usados: set = set()
    fecha_reporte = date.today()

    for i, bodega in enumerate(bodegas):
        hoja = libro.active if i == 0 else libro.create_sheet()
        hoja.title = _nombre_hoja_excel(bodega.nombre, nombres_usados)
        hoja.sheet_view.showGridLines = False

        # ---------- Encabezado de la hoja ----------
        hoja.merge_cells("A1:D1")
        hoja["A1"] = "REPORTE DE STOCK DISPONIBLE"
        hoja["A1"].font = Font(bold=True, size=16, color=azul)
        hoja["A1"].alignment = Alignment(horizontal="center", vertical="center")
        hoja.row_dimensions[1].height = 28

        encabezado = [
            ("Bodega:", bodega.nombre, None),
            ("Fecha del reporte:", fecha_reporte, "dd/mm/yyyy"),
            ("Generado por:", usuario.nombre, None),
        ]
        for fila, (etiqueta, valor, formato) in enumerate(encabezado, start=2):
            hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
            celda = hoja.cell(row=fila, column=2, value=valor)
            celda.alignment = Alignment(horizontal="left")
            if formato:
                celda.number_format = formato

        # ---------- Títulos de columnas ----------
        for col, (titulo, ancho, formato, alineacion) in enumerate(columnas, start=1):
            celda = hoja.cell(row=fila_titulos, column=col, value=titulo)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno_titulo
            celda.border = borde
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            hoja.column_dimensions[get_column_letter(col)].width = ancho
        hoja.row_dimensions[fila_titulos].height = 30

        # ---------- Datos ----------
        filas = filas_por_bodega.get(bodega.id, [])
        for n, r in enumerate(filas):
            num_fila = fila_titulos + 1 + n
            valores = [r.codigo, r.producto, r.presentacion, r.stock]
            for col, ((titulo, ancho, formato, alineacion), valor) in enumerate(zip(columnas, valores), start=1):
                celda = hoja.cell(row=num_fila, column=col, value=valor)
                celda.border = borde
                celda.alignment = Alignment(horizontal=alineacion, vertical="center")
                if formato:
                    celda.number_format = formato
                if n % 2 == 1:
                    celda.fill = relleno_alterno

        ultima_fila = fila_titulos + len(filas)
        if not filas:
            hoja.cell(row=fila_titulos + 1, column=1, value="Sin productos con stock en esta bodega").font = Font(italic=True)

        # ---------- Filtros, congelado e impresión ----------
        hoja.auto_filter.ref = f"A{fila_titulos}:D{max(ultima_fila, fila_titulos)}"
        hoja.freeze_panes = f"A{fila_titulos + 1}"
        hoja.print_title_rows = f"{fila_titulos}:{fila_titulos}"
        hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"stock_por_bodega_{fecha_reporte:%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
    
@router.get("/valor-inventario/excel")
def valor_inventario_excel(
    db: Session = Depends(get_db),
    usuario=Depends(require_permission("inventario.ver")),
):
    filas = db.execute(text("""
        SELECT codigo, producto, bodega, cantidad, costo_promedio, valor
        FROM public.v_valor_inventario
        ORDER BY bodega, producto
    """)).all()
 
    azul = "1F6FEB"
    relleno_titulo = PatternFill("solid", fgColor=azul)
    relleno_alterno = PatternFill("solid", fgColor="F3F7FF")
    relleno_total = PatternFill("solid", fgColor="E8F0FE")
    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
 
    columnas = [
        ("CÓDIGO", 20, None, "left"),
        ("PRODUCTO", 52, None, "left"),
        ("BODEGA", 24, None, "left"),
        ("CANTIDAD", 16, "#,##0.000", "right"),
        ("COSTO PROMEDIO", 18, "$#,##0.0000", "right"),
        ("VALOR TOTAL", 18, "$#,##0.00", "right"),
    ]
    fila_titulos = 6
    fecha_reporte = date.today()
 
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Valor de inventario"
    hoja.sheet_view.showGridLines = False
 
    # ---------- Encabezado de la hoja ----------
    hoja.merge_cells("A1:F1")
    hoja["A1"] = "REPORTE DE VALOR DE INVENTARIO"
    hoja["A1"].font = Font(bold=True, size=16, color=azul)
    hoja["A1"].alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[1].height = 28
 
    encabezado = [
        ("Fecha del reporte:", fecha_reporte, "dd/mm/yyyy"),
        ("Generado por:", usuario.nombre, None),
    ]
    for fila, (etiqueta, valor, formato) in enumerate(encabezado, start=2):
        hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda = hoja.cell(row=fila, column=2, value=valor)
        celda.alignment = Alignment(horizontal="left")
        if formato:
            celda.number_format = formato
 
    # ---------- Títulos de columnas ----------
    for col, (titulo, ancho, formato, alineacion) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_titulos, column=col, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno_titulo
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[fila_titulos].height = 30
 
    # ---------- Datos ----------
    for n, r in enumerate(filas):
        num_fila = fila_titulos + 1 + n
        valores = [r.codigo, r.producto, r.bodega, r.cantidad, r.costo_promedio, r.valor]
        for col, ((titulo, ancho, formato, alineacion), valor) in enumerate(zip(columnas, valores), start=1):
            celda = hoja.cell(row=num_fila, column=col, value=valor)
            celda.border = borde
            celda.alignment = Alignment(horizontal=alineacion, vertical="center")
            if formato:
                celda.number_format = formato
            if n % 2 == 1:
                celda.fill = relleno_alterno
 
    ultima_fila = fila_titulos + len(filas)
 
    if not filas:
        hoja.cell(row=fila_titulos + 1, column=1, value="Sin productos con valor en inventario").font = Font(italic=True)
    else:
        # ---------- Total general (se recalcula según los filtros aplicados) ----------
        fila_total = ultima_fila + 2
        hoja.merge_cells(start_row=fila_total, start_column=1, end_row=fila_total, end_column=5)
        etiqueta = hoja.cell(row=fila_total, column=1, value="VALOR TOTAL DEL INVENTARIO")
        etiqueta.font = Font(bold=True)
        etiqueta.alignment = Alignment(horizontal="right", vertical="center")
        total = hoja.cell(
            row=fila_total,
            column=6,
            value=f"=SUBTOTAL(109,F{fila_titulos + 1}:F{ultima_fila})",
        )
        total.font = Font(bold=True, color=azul)
        total.number_format = "$#,##0.00"
        total.alignment = Alignment(horizontal="right", vertical="center")
        for col in range(1, 7):
            hoja.cell(row=fila_total, column=col).fill = relleno_total
            hoja.cell(row=fila_total, column=col).border = borde
 
    # ---------- Filtros, congelado e impresión ----------
    hoja.auto_filter.ref = f"A{fila_titulos}:F{max(ultima_fila, fila_titulos)}"
    hoja.freeze_panes = f"A{fila_titulos + 1}"
    hoja.print_title_rows = f"{fila_titulos}:{fila_titulos}"
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
 
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
 
    nombre_archivo = f"valor_inventario_{fecha_reporte:%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
    
# =====================================================================
# NUEVO: Reporte de productos agotados (una sola hoja, con filtros)
# Ruta: GET /reportes/productos-agotados/excel
# Lee la vista v_productos_agotados
# =====================================================================
@router.get("/productos-agotados/excel")
def productos_agotados_excel(
    db: Session = Depends(get_db),
    usuario=Depends(require_permission("inventario.ver")),
):
    filas = db.execute(text("""
        SELECT codigo, producto, cantidad_total
        FROM public.v_productos_agotados
        ORDER BY producto
    """)).all()

    azul = "1F6FEB"
    relleno_titulo = PatternFill("solid", fgColor=azul)
    relleno_alterno = PatternFill("solid", fgColor="F3F7FF")
    relleno_total = PatternFill("solid", fgColor="E8F0FE")
    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    columnas = [
        ("CÓDIGO", 22, None, "left"),
        ("PRODUCTO", 60, None, "left"),
        ("CANTIDAD TOTAL", 18, "#,##0.000", "right"),
    ]
    fila_titulos = 6
    fecha_reporte = date.today()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Productos agotados"
    hoja.sheet_view.showGridLines = False

    # ---------- Encabezado de la hoja ----------
    hoja.merge_cells("A1:C1")
    hoja["A1"] = "REPORTE DE PRODUCTOS AGOTADOS"
    hoja["A1"].font = Font(bold=True, size=16, color=azul)
    hoja["A1"].alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[1].height = 28

    encabezado = [
        ("Fecha del reporte:", fecha_reporte, "dd/mm/yyyy"),
        ("Generado por:", usuario.nombre, None),
    ]
    for fila, (etiqueta, valor, formato) in enumerate(encabezado, start=2):
        hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda = hoja.cell(row=fila, column=2, value=valor)
        celda.alignment = Alignment(horizontal="left")
        if formato:
            celda.number_format = formato

    # ---------- Títulos de columnas ----------
    for col, (titulo, ancho, formato, alineacion) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_titulos, column=col, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno_titulo
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[fila_titulos].height = 30

    # ---------- Datos ----------
    for n, r in enumerate(filas):
        num_fila = fila_titulos + 1 + n
        valores = [r.codigo, r.producto, r.cantidad_total]
        for col, ((titulo, ancho, formato, alineacion), valor) in enumerate(zip(columnas, valores), start=1):
            celda = hoja.cell(row=num_fila, column=col, value=valor)
            celda.border = borde
            celda.alignment = Alignment(horizontal=alineacion, vertical="center")
            if formato:
                celda.number_format = formato
            if n % 2 == 1:
                celda.fill = relleno_alterno

    ultima_fila = fila_titulos + len(filas)

    if not filas:
        hoja.cell(row=fila_titulos + 1, column=1, value="No hay productos agotados").font = Font(italic=True)
    else:
        # ---------- Total (se recalcula según los filtros aplicados) ----------
        fila_total = ultima_fila + 2
        hoja.merge_cells(start_row=fila_total, start_column=1, end_row=fila_total, end_column=2)
        etiqueta = hoja.cell(row=fila_total, column=1, value="TOTAL DE PRODUCTOS AGOTADOS")
        etiqueta.font = Font(bold=True)
        etiqueta.alignment = Alignment(horizontal="right", vertical="center")
        total = hoja.cell(
            row=fila_total,
            column=3,
            value=f"=SUBTOTAL(103,A{fila_titulos + 1}:A{ultima_fila})",
        )
        total.font = Font(bold=True, color=azul)
        total.number_format = "0"
        total.alignment = Alignment(horizontal="right", vertical="center")
        for col in range(1, 4):
            hoja.cell(row=fila_total, column=col).fill = relleno_total
            hoja.cell(row=fila_total, column=col).border = borde

    # ---------- Filtros, congelado e impresión ----------
    hoja.auto_filter.ref = f"A{fila_titulos}:C{max(ultima_fila, fila_titulos)}"
    hoja.freeze_panes = f"A{fila_titulos + 1}"
    hoja.print_title_rows = f"{fila_titulos}:{fila_titulos}"
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"productos_agotados_{fecha_reporte:%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )    
    
    # =====================================================================
# NUEVO: Reporte de movimientos entre bodegas (una sola hoja, con filtros)
# Ruta: GET /reportes/movimientos-entre-bodegas/excel
# Lee la vista movimientos_entre_bodegas
# =====================================================================
def _sin_zona_horaria(fecha):
    # Excel no acepta fechas con zona horaria, así que se la quitamos
    if fecha is None:
        return None
    return fecha.replace(tzinfo=None)


@router.get("/movimientos-entre-bodegas/excel")
def movimientos_entre_bodegas_excel(
    db: Session = Depends(get_db),
    usuario=Depends(require_permission("inventario.ver")),
):
    filas = db.execute(text("""
        SELECT movimiento_id, fecha, estado, bodega_origen, bodega_destino,
               codigo, producto, cantidad, unidad_base, costo_unitario,
               valor_linea, fecha_recepcion, observacion_despacho,
               observacion_recepcion
        FROM public.movimientos_entre_bodegas
        ORDER BY fecha DESC, movimiento_id DESC
    """)).all()

    azul = "1F6FEB"
    relleno_titulo = PatternFill("solid", fgColor=azul)
    relleno_alterno = PatternFill("solid", fgColor="F3F7FF")
    relleno_total = PatternFill("solid", fgColor="E8F0FE")
    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    # (título, ancho, formato, alineación, ajustar texto)
    columnas = [
        ("MOVIMIENTO", 14, "0", "center", False),
        ("FECHA", 18, "dd/mm/yyyy hh:mm", "center", False),
        ("ESTADO", 15, None, "center", False),
        ("BODEGA ORIGEN", 20, None, "left", False),
        ("BODEGA DESTINO", 20, None, "left", False),
        ("CÓDIGO", 18, None, "left", False),
        ("PRODUCTO", 48, None, "left", False),
        ("CANTIDAD", 13, "#,##0.000", "right", False),
        ("UNIDAD", 10, None, "center", False),
        ("COSTO UNITARIO", 17, "$#,##0.0000", "right", False),
        ("VALOR", 14, "$#,##0.00", "right", False),
        ("FECHA RECEPCIÓN", 18, "dd/mm/yyyy hh:mm", "center", False),
        ("OBSERVACIÓN DESPACHO", 36, None, "left", True),
        ("OBSERVACIÓN RECEPCIÓN", 36, None, "left", True),
    ]
    fila_titulos = 6
    ultima_columna = get_column_letter(len(columnas))
    fecha_reporte = date.today()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Movimientos entre bodegas"
    hoja.sheet_view.showGridLines = False

    # ---------- Encabezado de la hoja ----------
    hoja.merge_cells(f"A1:{ultima_columna}1")
    hoja["A1"] = "REPORTE DE MOVIMIENTOS ENTRE BODEGAS"
    hoja["A1"].font = Font(bold=True, size=16, color=azul)
    hoja["A1"].alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[1].height = 28

    encabezado = [
        ("Fecha del reporte:", fecha_reporte, "dd/mm/yyyy"),
        ("Generado por:", usuario.nombre, None),
    ]
    for fila, (etiqueta, valor, formato) in enumerate(encabezado, start=2):
        hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda = hoja.cell(row=fila, column=2, value=valor)
        celda.alignment = Alignment(horizontal="left")
        if formato:
            celda.number_format = formato

    # ---------- Títulos de columnas ----------
    for col, (titulo, ancho, formato, alineacion, ajustar) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_titulos, column=col, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno_titulo
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[fila_titulos].height = 30

    # ---------- Datos ----------
    for n, r in enumerate(filas):
        num_fila = fila_titulos + 1 + n
        valores = [
            r.movimiento_id,
            _sin_zona_horaria(r.fecha),
            r.estado,
            r.bodega_origen,
            r.bodega_destino,
            r.codigo,
            r.producto,
            r.cantidad,
            r.unidad_base,
            r.costo_unitario,
            r.valor_linea,
            _sin_zona_horaria(r.fecha_recepcion),
            r.observacion_despacho,
            r.observacion_recepcion,
        ]
        for col, ((titulo, ancho, formato, alineacion, ajustar), valor) in enumerate(zip(columnas, valores), start=1):
            celda = hoja.cell(row=num_fila, column=col, value=valor)
            celda.border = borde
            celda.alignment = Alignment(horizontal=alineacion, vertical="center", wrap_text=ajustar)
            if formato:
                celda.number_format = formato
            if n % 2 == 1:
                celda.fill = relleno_alterno

    ultima_fila = fila_titulos + len(filas)

    if not filas:
        hoja.cell(row=fila_titulos + 1, column=1, value="No hay movimientos entre bodegas").font = Font(italic=True)
    else:
        # ---------- Total de valor (se recalcula según los filtros aplicados) ----------
        fila_total = ultima_fila + 2
        columna_valor = 11  # columna K: VALOR
        hoja.merge_cells(start_row=fila_total, start_column=1, end_row=fila_total, end_column=columna_valor - 1)
        etiqueta = hoja.cell(row=fila_total, column=1, value="VALOR TOTAL DE LOS MOVIMIENTOS")
        etiqueta.font = Font(bold=True)
        etiqueta.alignment = Alignment(horizontal="right", vertical="center")
        letra_valor = get_column_letter(columna_valor)
        total = hoja.cell(
            row=fila_total,
            column=columna_valor,
            value=f"=SUBTOTAL(109,{letra_valor}{fila_titulos + 1}:{letra_valor}{ultima_fila})",
        )
        total.font = Font(bold=True, color=azul)
        total.number_format = "$#,##0.00"
        total.alignment = Alignment(horizontal="right", vertical="center")
        for col in range(1, columna_valor + 1):
            hoja.cell(row=fila_total, column=col).fill = relleno_total
            hoja.cell(row=fila_total, column=col).border = borde

    # ---------- Filtros, congelado e impresión ----------
    hoja.auto_filter.ref = f"A{fila_titulos}:{ultima_columna}{max(ultima_fila, fila_titulos)}"
    hoja.freeze_panes = f"A{fila_titulos + 1}"
    hoja.print_title_rows = f"{fila_titulos}:{fila_titulos}"
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"movimientos_entre_bodegas_{fecha_reporte:%Y-%m-%d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )